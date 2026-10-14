"""Reportes XLSX de subredes: individual y general (todas las subredes en una hoja)."""
import io
import ipaddress
from math import ceil

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.Models.ip_address import IPStatus

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
ROWS_PER_COLUMN = 52
MAX_HOSTS = 4096

# (fondo, texto): mismos colores del dashboard/SPA.
COLORS = {
    IPStatus.FREE: ("198754", "FFFFFF"),
    IPStatus.ASSIGNED: ("DC3545", "FFFFFF"),
    IPStatus.ACTIVE: ("FFC107", "212529"),
}
LABELS = {
    IPStatus.FREE: "Libre",
    IPStatus.ASSIGNED: "Asignada",
    IPStatus.ACTIVE: "Activa sin registrar",
}

_THIN = Side(style="thin", color="D9D9D9")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_CENTER = Alignment(horizontal="center", vertical="center")


class SubnetReportError(ValueError):
    pass


def _host_label(ip: ipaddress.IPv4Address, prefix: int):
    k = 1 if prefix >= 24 else 2 if prefix >= 16 else 3
    parts = str(ip).split(".")[-k:]
    return int(parts[0]) if k == 1 else ".".join(parts)


def _parse_network(subnet):
    """Valida la subred y devuelve su ip_network; lanza SubnetReportError si no es exportable."""
    try:
        net = ipaddress.ip_network(subnet.cidr, strict=False)
    except ValueError as exc:
        raise SubnetReportError(f"CIDR inválido: {subnet.cidr}") from exc
    if net.version != 4:
        raise SubnetReportError("La exportación solo soporta subredes IPv4.")
    if net.num_addresses > MAX_HOSTS + 2:
        raise SubnetReportError(
            f"La subred tiene {net.num_addresses} direcciones; máximo exportable: {MAX_HOSTS}."
        )
    return net


def _draw_block(ws, subnet, net, ips, start_col: int) -> int:
    """Dibuja el bloque de una subred desde `start_col`. Devuelve las columnas usadas."""
    hosts = list(net.hosts())
    status_by_ip = {i.ip_address: i.status for i in ips}
    ncols = max(1, ceil(len(hosts) / ROWS_PER_COLUMN))

    # Header: subred completa
    ws.merge_cells(start_row=1, start_column=start_col, end_row=1, end_column=start_col + ncols - 1)
    head = ws.cell(row=1, column=start_col, value=f"{subnet.cidr} {subnet.name}")
    head.font = Font(bold=True, color="FFFFFF", size=12)
    head.fill = PatternFill("solid", fgColor="1F4E78")
    head.alignment = _CENTER

    for c in range(start_col, start_col + ncols):
        ws.column_dimensions[get_column_letter(c)].width = 12

    # Hosts: llenado por columnas (52 filas c/u)
    for idx, host in enumerate(hosts):
        cell = ws.cell(row=2 + idx % ROWS_PER_COLUMN,
                       column=start_col + idx // ROWS_PER_COLUMN,
                       value=_host_label(host, net.prefixlen))
        cell.alignment = _CENTER
        cell.border = _BORDER
        status = status_by_ip.get(str(host))
        if status in COLORS:
            bg, fg = COLORS[status]
            cell.fill = PatternFill("solid", fgColor=bg)
            cell.font = Font(color=fg)
    return ncols


def _draw_legend(ws, sw: int, skipped: list[str] | None = None) -> None:
    """Leyenda de colores: swatch en columna `sw`, texto en `sw + 1`."""
    tx = sw + 1
    ws.column_dimensions[get_column_letter(sw)].width = 4
    ws.column_dimensions[get_column_letter(tx)].width = 30
    for n, st in enumerate(IPStatus):
        bg, _ = COLORS[st]
        ws.cell(row=2 + n, column=sw).fill = PatternFill("solid", fgColor=bg)
        ws.cell(row=2 + n, column=tx, value=LABELS[st])
    row = 2 + len(IPStatus)
    ws.cell(row=row, column=tx, value="Sin color: no registrada en el sistema")
    for msg in skipped or []:
        row += 1
        ws.cell(row=row, column=tx, value=msg).font = Font(italic=True, color="DC3545")


def _to_bytes(wb: Workbook) -> bytes:
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def build_subnet_report(subnet, ips) -> bytes:
    net = _parse_network(subnet)
    wb = Workbook()
    ws = wb.active
    ws.title = "Reporte Subred"
    ws.row_dimensions[1].height = 22
    ncols = _draw_block(ws, subnet, net, ips, start_col=1)
    _draw_legend(ws, sw=ncols + 2)
    return _to_bytes(wb)


def build_general_report(data: list[tuple]) -> bytes:
    """`data` = [(subnet, ips), ...]. Una hoja, bloques lado a lado, una sola leyenda.
    Las subredes no exportables (IPv6, CIDR inválido, muy grandes) se omiten y se avisan en la leyenda."""
    if not data:
        raise SubnetReportError("No hay subredes para exportar.")

    wb = Workbook()
    ws = wb.active
    ws.title = "Reporte General"
    ws.row_dimensions[1].height = 22

    col, skipped = 1, []
    for subnet, ips in data:
        try:
            net = _parse_network(subnet)
        except SubnetReportError:
            skipped.append(f"Omitida: {subnet.cidr}")
            continue
        ncols = _draw_block(ws, subnet, net, ips, start_col=col)
        col += ncols + 1  # + columna separadora
        ws.column_dimensions[get_column_letter(col - 1)].width = 3

    if col == 1:
        raise SubnetReportError("Ninguna subred es exportable (solo IPv4 hasta 4096 hosts).")

    _draw_legend(ws, sw=col, skipped=skipped)
    return _to_bytes(wb)